#!/usr/bin/env python3
"""Carga email_ejecutivo desde BASE PARA PRUEBA.xlsx → mnp_cartera_oficial.

Instalar dependencias: python3 -m pip install -r scripts/requirements.txt
"""
from openpyxl import load_workbook
import os
import sys

try:
    import psycopg2
except ImportError:
    psycopg2 = None  # type: ignore[assignment]

EXCEL = os.environ.get(
    "CARTERA_XLSX",
    "/opt/mol-docker/docs/becas-beneficios/BASE PARA PRUEBA.xlsx",
)


def rut_norm(rut, dig):
    body = "".join(c for c in str(rut) if c.isdigit())
    dv = str(dig).strip().upper()
    return body + dv


def load_emails_by_rut():
    wb = load_workbook(EXCEL, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    idx = {h.strip(): i for i, h in enumerate(header)}
    by_rut = {}
    for row in rows:
        if row is None or row[idx["RUT"]] is None:
            continue
        rn = rut_norm(row[idx["RUT"]], row[idx["DIG"]])
        email = (row[idx["EJECUTIVO MATRICULA"]] or "").strip().lower()
        if not email:
            continue
        by_rut.setdefault(rn, email)
    wb.close()
    return by_rut


def get_dsn():
    dsn = os.environ.get("SUPABASE_PG_URL") or os.environ.get("SUPABASE_DATABASE_URL")
    if dsn:
        return dsn
    host = os.environ.get("SUPABASE_PG_HOST")
    if not host:
        return None
    port = os.environ.get("SUPABASE_PG_PORT", "5432")
    db = os.environ.get("SUPABASE_PG_DB", "postgres")
    user = os.environ.get("SUPABASE_PG_USER", "postgres")
    password = os.environ.get("SUPABASE_PG_PASSWORD", "")
    return f"postgresql://{user}:{password}@{host}:{port}/{db}"


def main():
    by_rut = load_emails_by_rut()
    distinct_emails = sorted(set(by_rut.values()))
    print(f"excel_rows={len(by_rut)} distinct_emails={len(distinct_emails)}")
    for email in distinct_emails:
        print(f"  {email}")

    dsn = get_dsn()
    if not dsn:
        print("SUPABASE_PG_URL not set; skipping DB update", file=sys.stderr)
        sys.exit(0 if by_rut else 1)

    if psycopg2 is None:
        print("psycopg2 not installed; skipping DB update", file=sys.stderr)
        sys.exit(0 if by_rut else 1)

    conn = psycopg2.connect(dsn)
    cur = conn.cursor()
    n = 0
    for rn, email in by_rut.items():
        cur.execute(
            "UPDATE public.mnp_cartera_oficial SET email_ejecutivo = %s WHERE rut_norm = %s",
            (email, rn),
        )
        n += cur.rowcount
    conn.commit()
    cur.close()
    conn.close()
    print(f"updated={n} distinct_rut={len(by_rut)}")


if __name__ == "__main__":
    main()
