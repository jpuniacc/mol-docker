#!/usr/bin/env python3
"""Carga becas desde BASE PARA PRUEBA.xlsx → mnp_cartera_beneficios.

Instalar dependencias: python3 -m pip install -r scripts/requirements.txt
"""
from __future__ import annotations

import csv
import os
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

try:
    import psycopg2
except ImportError:
    psycopg2 = None  # type: ignore[assignment]

EXCEL = os.environ.get(
    "CARTERA_XLSX",
    "/opt/mol-docker/docs/becas-beneficios/BASE PARA PRUEBA.xlsx",
)
PERIODO = os.environ.get("CARTERA_PERIODO", "2027-01")
CATALOG_CSV = os.environ.get(
    "BENEFICIOS_CATALOG_CSV",
    "/opt/mol-docker/docs/becas-beneficios/cod_beneficios_flujos_2027-01.csv",
)
CONSOLIDADO_COL = "CONSOLIDADO CAE-BECA MINISTERIAL-SUBDERE"


def normalizar_nombre_beca(nombre: str | None) -> str:
    text = unicodedata.normalize("NFKC", nombre or "")
    return " ".join(text.strip().split()).lower()


def es_sin_beca(nombre: str | None) -> bool:
    n = normalizar_nombre_beca(nombre)
    return not n or n in {"sin beca", "-", "n/a"}


def match_cod_beneficio(
    nombre_excel: str | None,
    catalogo: list[tuple[str, str]],
) -> str | None:
    if es_sin_beca(nombre_excel):
        return None
    target = normalizar_nombre_beca(nombre_excel)
    for codigo, beneficio in catalogo:
        if normalizar_nombre_beca(beneficio) == target:
            return codigo.strip() or None
    return None


def rut_norm(rut: Any, dig: Any) -> str:
    body = "".join(c for c in str(rut) if c.isdigit())
    dv = str(dig).strip().upper()
    return body + dv


def parse_pct(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).strip()
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def cell_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def get_dsn() -> str | None:
    dsn = os.environ.get("SUPABASE_PG_URL") or os.environ.get("SUPABASE_DATABASE_URL")
    if dsn:
        return dsn
    host = os.environ.get("SUPABASE_PG_HOST")
    if not host:
        return None
    port = os.environ.get("SUPABASE_PG_PORT", "5432")
    db = os.environ.get("SUPABASE_PG_DB", "postgres")
    user = os.environ.get("SUPABASE_PG_USER", "postgres")
    password = os.environ.get("SUPABASE_PG_PASSWORD", "")
    return f"postgresql://{user}:{password}@{host}:{port}/{db}"


def load_catalog_from_csv() -> list[tuple[str, str]]:
    catalog: list[tuple[str, str]] = []
    with open(CATALOG_CSV, newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        next(reader, None)
        for row in reader:
            if len(row) >= 2 and row[0].strip():
                catalog.append((row[0].strip(), row[1].strip()))
    return catalog


def load_catalog_from_db(cur) -> list[tuple[str, str]]:
    cur.execute(
        """
        SELECT codigo_beneficio::text, beneficio
        FROM public.mnp_mv_beneficio_periodo
        WHERE periodo = %s
        ORDER BY codigo_beneficio
        """,
        (PERIODO,),
    )
    rows = cur.fetchall()
    return [(str(cod), str(nom)) for cod, nom in rows]


def resolve_column_indices(header: list[str]) -> dict[str, int]:
    idx = {h.strip(): i for i, h in enumerate(header)}
    required = ["CODCLI", "RUT", "DIG", "CODCARPR", "BECA 1", "BECA 2"]
    missing = [name for name in required if name not in idx]
    if missing:
        raise SystemExit(f"Columnas faltantes en Excel: {', '.join(missing)}")
    beca1_i = idx["BECA 1"]
    beca2_i = idx["BECA 2"]
    return {
        "codcli": idx["CODCLI"],
        "rut": idx["RUT"],
        "dig": idx["DIG"],
        "codcarpr": idx["CODCARPR"],
        "beca1": beca1_i,
        "pct1": beca1_i + 1,
        "beca2": beca2_i,
        "pct2": beca2_i + 1,
        "consolidado": idx.get(CONSOLIDADO_COL),
    }


def load_excel_rows() -> list[dict[str, Any]]:
    wb = load_workbook(EXCEL, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows_iter)]
    cols = resolve_column_indices(header)
    out: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or row[cols["codcli"]] is None:
            continue
        codcli = str(row[cols["codcli"]]).strip()
        if not codcli:
            continue
        beca1 = cell_text(row[cols["beca1"]])
        beca2 = cell_text(row[cols["beca2"]])
        consolidado = (
            cell_text(row[cols["consolidado"]]) if cols["consolidado"] is not None else None
        )
        out.append(
            {
                "codcli_excel": codcli,
                "rut_norm": rut_norm(row[cols["rut"]], row[cols["dig"]]),
                "codcarpr": cell_text(row[cols["codcarpr"]]),
                "beca_1": None if es_sin_beca(beca1) else beca1,
                "pct_1": parse_pct(row[cols["pct1"]]),
                "beca_2": None if es_sin_beca(beca2) else beca2,
                "pct_2": parse_pct(row[cols["pct2"]]),
                "consolidado": consolidado,
            }
        )
    wb.close()
    return out


UPSERT_SQL = """
INSERT INTO public.mnp_cartera_beneficios (
    periodo, rut_norm, codcli_excel, codcarpr,
    beca_1, pct_1, beca_2, pct_2, consolidado,
    cod_beneficio_1, cod_beneficio_2, loaded_at
) VALUES (
    %s, %s, %s, %s,
    %s, %s, %s, %s, %s,
    %s, %s, now()
)
ON CONFLICT (periodo, codcli_excel) DO UPDATE SET
    rut_norm = EXCLUDED.rut_norm,
    codcarpr = EXCLUDED.codcarpr,
    beca_1 = EXCLUDED.beca_1,
    pct_1 = EXCLUDED.pct_1,
    beca_2 = EXCLUDED.beca_2,
    pct_2 = EXCLUDED.pct_2,
    consolidado = EXCLUDED.consolidado,
    cod_beneficio_1 = EXCLUDED.cod_beneficio_1,
    cod_beneficio_2 = EXCLUDED.cod_beneficio_2,
    loaded_at = now()
"""


def enrich_and_count(rows: list[dict[str, Any]], catalogo: list[tuple[str, str]]) -> tuple[int, int, int, int]:
    matched1 = unmatched1 = matched2 = unmatched2 = 0
    for row in rows:
        row["cod_beneficio_1"] = match_cod_beneficio(row["beca_1"], catalogo)
        row["cod_beneficio_2"] = match_cod_beneficio(row["beca_2"], catalogo)
        if row["beca_1"]:
            if row["cod_beneficio_1"]:
                matched1 += 1
            else:
                unmatched1 += 1
        if row["beca_2"]:
            if row["cod_beneficio_2"]:
                matched2 += 1
            else:
                unmatched2 += 1
    return matched1, unmatched1, matched2, unmatched2


def upsert_rows(cur, rows: list[dict[str, Any]]) -> int:
    n = 0
    for row in rows:
        cur.execute(
            UPSERT_SQL,
            (
                PERIODO,
                row["rut_norm"],
                row["codcli_excel"],
                row["codcarpr"],
                row["beca_1"],
                row["pct_1"],
                row["beca_2"],
                row["pct_2"],
                row["consolidado"],
                row["cod_beneficio_1"],
                row["cod_beneficio_2"],
            ),
        )
        n += 1
    return n


def main() -> None:
    rows = load_excel_rows()
    dsn = get_dsn()
    catalog_source = "csv"
    catalogo = load_catalog_from_csv()

    if dsn and psycopg2 is not None:
        conn = psycopg2.connect(dsn)
        cur = conn.cursor()
        db_catalog = load_catalog_from_db(cur)
        if db_catalog:
            catalogo = db_catalog
            catalog_source = "db"
        matched1, unmatched1, matched2, unmatched2 = enrich_and_count(rows, catalogo)
        upserted = upsert_rows(cur, rows)
        conn.commit()
        cur.execute(
            """
            SELECT beca_1, cod_beneficio_1, beca_2, cod_beneficio_2
            FROM public.mnp_cartera_beneficios
            WHERE periodo = %s AND codcli_excel = %s
            """,
            (PERIODO, "20152PSIC1SR039"),
        )
        vanessa = cur.fetchone()
        cur.close()
        conn.close()
        print(
            f"rows={len(rows)} matched1={matched1} unmatched1={unmatched1} "
            f"matched2={matched2} unmatched2={unmatched2} upserted={upserted} "
            f"catalog={catalog_source}({len(catalogo)})"
        )
        if vanessa:
            print(f"vanessa={vanessa}")
        else:
            print("vanessa=NOT_FOUND", file=sys.stderr)
            sys.exit(1)
        if vanessa[1] != "1756":
            print(f"vanessa cod_beneficio_1 expected 1756 got {vanessa[1]}", file=sys.stderr)
            sys.exit(1)
        return

    matched1, unmatched1, matched2, unmatched2 = enrich_and_count(rows, catalogo)
    print(
        f"rows={len(rows)} matched1={matched1} unmatched1={unmatched1} "
        f"matched2={matched2} unmatched2={unmatched2} catalog={catalog_source}({len(catalogo)})"
    )
    if not dsn:
        print("SUPABASE_PG_URL not set; skipping DB upsert", file=sys.stderr)
    elif psycopg2 is None:
        print("psycopg2 not installed; skipping DB upsert", file=sys.stderr)
    sys.exit(0 if rows else 1)


if __name__ == "__main__":
    main()
